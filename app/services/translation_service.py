from pathlib import Path
from typing import Callable, Dict, Optional

from app.core import constants
from app.core.exceptions import ExtractionError, ReconstructionError
from app.core.logging_config import logger
from app.core.security import result_path_for, sha256_of_file
from app.interfaces.formatter import ExtractionResult
from app.services.correction.linguistic_corrector import LinguisticCorrector
from app.services.formatters.registry import registry
from app.services.glossary_service import GlossaryService
from app.services.protection.element_protector import ElementProtector
from app.services.task_service import TaskService
from app.services.translation.engine_manager import engine_manager

ProgressCb = Callable[[int, int, str], None]


class TranslationPipeline:
    """End-to-end orchestration of a translation job.

    Steps: validate -> extract -> detect language -> protect -> translate ->
    (optional) correct -> reconstruct -> verify. Supports resume by skipping
    segments already translated and persisted.
    """

    def __init__(self, db, progress: Optional[ProgressCb] = None, resume: bool = False):
        self.db = db
        self.tasks = TaskService(db)
        self.glossary = GlossaryService(db)
        self.progress = progress or (lambda done, total, msg: None)
        self.resume = resume

    def run(self, task_id: str, allow_auto_download: bool = False) -> Dict:
        task = self.tasks.get(task_id)
        if not task:
            raise ValueError(f"Tâche inconnue : {task_id}")

        source_path = Path(task.source_path)
        self.tasks.set_status(task_id, constants.TASK_IN_PROGRESS)

        try:
            # 1. Verify source integrity on resume.
            if task.source_checksum:
                current = sha256_of_file(source_path)
                if current != task.source_checksum:
                    raise ExtractionError(
                        "Le fichier source a changé depuis la création de la tâche. "
                        "Reprise impossible."
                    )

            # 2. Extract.
            self.progress(0, 1, "Extraction du document…")
            formatter = registry.get_for_path(source_path)
            options = self._load_options(task)
            extraction: ExtractionResult = formatter.extract(source_path, options)
            segments = extraction.segments
            total = len(segments)

            # 3. Language detection.
            if not task.source_lang and not task.detected_lang and segments:
                from app.services.language_detector import detector
                sample = "\n".join(s.text for s in segments[:20])
                detected = detector.detect(sample)
                task = self.tasks.get(task_id)
                task.detected_lang = detected
                self.tasks.db.commit()
            source_lang = task.source_lang or task.detected_lang

            # 4. Engine selection (local by default, remote requires consent).
            consent = bool(self._load_options(task).get("remote_consent", False))
            engine = engine_manager.select(
                requested=task.engine_used or None,
                consent=consent,
                source_lang=source_lang,
                target_lang=task.target_lang,
            )
            task = self.tasks.get(task_id)
            task.engine_used = engine.name
            task.engine_mode = engine.mode
            self.tasks.db.commit()

            # 5. Protection setup.
            protector = ElementProtector(
                custom_protected=self.glossary.protected_terms(),
                glossary={e.source_term: e.target_term
                          for e in self.glossary.list_entries()},
            )
            corrector = (LinguisticCorrector(language=task.target_lang)
                         if task.correction_enabled else None)

            # 6. Resume support.
            done_map: Dict[int, str] = {}
            if self.resume:
                done_map = self.tasks.load_translations(task_id)

            translations: Dict[int, str] = dict(done_map)
            processed = len(done_map)
            if total:
                self.tasks.update_progress(task_id, processed, total)

            for seg in segments:
                if seg.index in done_map:
                    continue
                # Context-aware, protection-safe translation.
                restored = protector.translate(
                    engine, seg.text, source_lang, task.target_lang)

                final = restored
                if corrector:
                    corrected, changes = corrector.correct(restored)
                    if changes and not corrector.is_meaning_preserving(restored, corrected):
                        # Keep the previous version if correction is too invasive.
                        corrected = restored
                    final = corrected

                # Engines usually trim edge whitespace, which is significant for
                # inline structures (HTML text nodes, inline markup, table cells).
                final = self._match_edge_whitespace(seg.text, final)

                translations[seg.index] = final
                self.tasks.save_translation(task_id, seg.index, restored, final,
                                            source_text=seg.text)
                processed += 1
                self.tasks.update_progress(task_id, processed, total)
                self.progress(processed, total,
                              f"Traduction {processed}/{total}")

            # 7. Reconstruct.
            self.progress(total, total, "Reconstruction du document…")
            output_path = self._output_path(task)
            rebuilt = formatter.reconstruct(source_path, translations,
                                            extraction.structure, output_path)

            # 8. Verify.
            self.progress(total, total, "Vérification du fichier généré…")
            self._verify_output(rebuilt, source_path)

            warnings = list(extraction.warnings)
            overflow = extraction.structure.get("_overflow_warnings")
            if overflow:
                warnings.extend(overflow)
            self.tasks.set_warnings(task_id, warnings)

            self.tasks.set_status(task_id, constants.TASK_COMPLETED,
                                  result_path=str(rebuilt))
            self.progress(total, total, "Terminé")
            return {
                "task_id": task_id,
                "status": constants.TASK_COMPLETED,
                "result_path": str(rebuilt),
                "segments": total,
                "engine": engine.name,
                "engine_mode": engine.mode,
                "warnings": warnings,
            }

        except Exception as exc:
            logger.exception("Échec de la tâche %s", task_id)
            self.tasks.add_error(task_id, str(exc))
            self.tasks.set_status(task_id, constants.TASK_FAILED)
            raise

    def _load_options(self, task) -> dict:
        from app.services.config_service import config_service
        return {
            "remote_consent": bool(config_service.get("remote_consent", False)),
        }

    @staticmethod
    def _match_edge_whitespace(source: str, translated: str) -> str:
        """Re-attach significant leading/trailing whitespace lost by engines."""
        if not translated or not source or not source.strip():
            return translated
        lead = source[:len(source) - len(source.lstrip())]
        trail = source[len(source.rstrip()):]
        if lead and not translated.startswith(lead):
            translated = lead + translated.lstrip()
        if trail and not translated.endswith(trail):
            translated = translated.rstrip() + trail
        return translated

    @staticmethod
    def _output_path(task) -> Path:
        return result_path_for(task.task_id, task.filename)

    @staticmethod
    def _verify_output(rebuilt: Path, source_path: Path) -> None:
        if not rebuilt.exists():
            raise ReconstructionError("Le fichier traduit n'a pas été généré.")
        if rebuilt.stat().st_size == 0:
            raise ReconstructionError("Le fichier traduit est vide.")
        ext = source_path.suffix.lower().lstrip(".")
        # Format-specific integrity checks where libraries allow.
        try:
            if ext == "docx":
                import docx
                docx.Document(str(rebuilt))
            elif ext == "pptx":
                from pptx import Presentation
                Presentation(str(rebuilt))
            elif ext == "xlsx":
                import openpyxl
                openpyxl.load_workbook(str(rebuilt))
            elif ext == "pdf":
                try:
                    import pymupdf as fitz
                except ImportError:
                    import fitz
                d = fitz.open(str(rebuilt))
                _ = d.page_count
                d.close()
            elif ext == "epub":
                import zipfile
                with zipfile.ZipFile(rebuilt, "r") as z:
                    if z.testzip() is not None:
                        raise ReconstructionError("Archive EPUB corrompue.")
        except ReconstructionError:
            raise
        except Exception as exc:
            raise ReconstructionError(f"Vérification du fichier échouée : {exc}")
