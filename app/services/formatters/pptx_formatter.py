from pathlib import Path
from typing import List

from app.interfaces.formatter import ExtractionResult, Segment
from app.services.formatters.base import BaseFormatter

from pptx import Presentation


class PptxFormatter(BaseFormatter):
    """PPTX: translate text frames, titles, tables and speaker notes.

    Images, layouts, slide order and non-text shapes are preserved. Text that
    risks overflowing its placeholder is reported as a warning.
    """

    extensions = (".pptx",)

    def extract(self, file_path: Path, options: dict = None) -> ExtractionResult:
        prs = Presentation(str(file_path))
        segments: List[Segment] = []
        warnings: List[str] = []
        idx = 0

        for si, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for pi, para in enumerate(shape.text_frame.paragraphs):
                        text = "".join(r.text for r in para.runs) or para.text
                        if text.strip():
                            segments.append(Segment(
                                index=idx, text=text,
                                ref=("shape", si, shape.shape_id, pi)))
                            idx += 1
                if shape.has_table:
                    for ri, row in enumerate(shape.table.rows):
                        for ci, cell in enumerate(row.cells):
                            if cell.text.strip():
                                segments.append(Segment(
                                    index=idx, text=cell.text,
                                    ref=("cell", si, shape.shape_id, ri, ci)))
                                idx += 1
            # Speaker notes.
            if slide.has_notes_slide:
                for pi, para in enumerate(slide.notes_slide.notes_text_frame.paragraphs):
                    text = para.text
                    if text.strip():
                        segments.append(Segment(
                            index=idx, text=text,
                            ref=("notes", si, pi)))
                        idx += 1

        return ExtractionResult(
            segments=segments,
            structure={"segments": segments},
            warnings=warnings,
        )

    def reconstruct(self, original_path, translations, structure, output_path):
        prs = Presentation(str(original_path))
        slides = list(prs.slides)
        shape_index = {}
        for si, slide in enumerate(slides):
            for shape in slide.shapes:
                shape_index[(si, shape.shape_id)] = shape

        overflow_warnings = []
        for seg in structure.get("segments", []):
            tr = translations.get(seg.index)
            if tr is None:
                continue
            ref = seg.ref
            kind = ref[0]
            if kind == "shape":
                _, si, sid, pi = ref
                shape = shape_index.get((si, sid))
                if not shape or not shape.has_text_frame:
                    continue
                if pi < len(shape.text_frame.paragraphs):
                    self._set_paragraph(shape.text_frame.paragraphs[pi], tr)
                    if self._likely_overflow(shape, seg.text, tr):
                        overflow_warnings.append(
                            f"Diapositive {si + 1}: texte possiblement trop long.")
            elif kind == "cell":
                _, si, sid, ri, ci = ref
                shape = shape_index.get((si, sid))
                if not shape or not shape.has_table:
                    continue
                table = shape.table
                if ri < len(table.rows) and ci < len(table.columns):
                    cell = table.cell(ri, ci)
                    self._set_cell(cell, tr)
                    if self._likely_overflow(cell, seg.text, tr):
                        overflow_warnings.append(
                            f"Diapositive {si + 1}: cellule possiblement trop longue.")
            elif kind == "notes":
                _, si, pi = ref
                if si < len(slides) and slides[si].has_notes_slide:
                    paras = slides[si].notes_slide.notes_text_frame.paragraphs
                    if pi < len(paras):
                        self._set_paragraph(paras[pi], tr)

        prs.save(str(output_path))
        structure["_overflow_warnings"] = list(dict.fromkeys(overflow_warnings))
        return output_path

    @staticmethod
    def _set_paragraph(para, text: str) -> None:
        runs = para.runs
        if not runs:
            para.add_run(text)
            return
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""

    @staticmethod
    def _set_cell(cell, text: str) -> None:
        tf = cell.text_frame
        if not tf.paragraphs:
            cell.text = text
            return
        para = tf.paragraphs[0]
        runs = para.runs
        if not runs:
            para.add_run(text)
        else:
            runs[0].text = text
            for run in runs[1:]:
                run.text = ""
        for extra in tf.paragraphs[1:]:
            for run in extra.runs:
                run.text = ""

    @staticmethod
    def _likely_overflow(shape_or_cell, source: str, translated: str) -> bool:
        try:
            width = None
            if hasattr(shape_or_cell, "width"):
                width = shape_or_cell.width
            if width:
                # Rough heuristic: chars per Emu width.
                capacity = max(1, int(width / 914400 * 12))
                return len(translated) > capacity * 4 and len(translated) > len(source) * 1.6
        except Exception:
            pass
        return False
