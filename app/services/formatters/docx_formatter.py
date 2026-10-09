from pathlib import Path
from typing import Iterator, List

from app.interfaces.formatter import ExtractionResult, Segment
from app.services.formatters.base import BaseFormatter

import docx
from docx.document import Document as _Doc
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


def iter_block_items(parent) -> Iterator:
    """Yield paragraphs and tables in document order."""
    if isinstance(parent, _Doc):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._element
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _tables_paragraphs(document) -> Iterator[Paragraph]:
    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            yield block
        elif isinstance(block, Table):
            for row in block.rows:
                seen_cells = set()
                for cell in row.cells:
                    if id(cell._tc) in seen_cells:
                        continue  # merged cell already yielded
                    seen_cells.add(id(cell._tc))
                    for para in cell.paragraphs:
                        yield para


def _container_paragraphs(container) -> Iterator[Paragraph]:
    try:
        for para in container.paragraphs:
            yield para
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        yield para
    except Exception:
        return


def iter_all_paragraphs(document) -> Iterator[Paragraph]:
    """Yields all translatable paragraphs in a deterministic order.

    Used identically at extraction and reconstruction so segment refs line up.
    """
    for para in _tables_paragraphs(document):
        yield para
    for section in document.sections:
        for container in (section.header, section.footer,
                          section.first_page_header, section.first_page_footer,
                          section.even_page_header, section.even_page_footer):
            if container is None:
                continue
            for para in _container_paragraphs(container):
                yield para


class DocxFormatter(BaseFormatter):
    """DOCX: preserve paragraphs, styles, tables, headers and footers."""

    extensions = (".docx",)

    def extract(self, file_path: Path, options: dict = None) -> ExtractionResult:
        document = docx.Document(str(file_path))
        segments: List[Segment] = []
        warnings: List[str] = []
        idx = 0
        for para in iter_all_paragraphs(document):
            text = para.text
            if not text.strip():
                continue
            runs_with_text = [r for r in para.runs if r.text]
            if len(runs_with_text) > 1 and len({(r.bold, r.italic) for r in runs_with_text}) > 1:
                warnings.append(
                    "Une mise en forme mixte (gras/italique) dans un paragraphe "
                    "peut être simplifiée après traduction."
                )
            segments.append(Segment(index=idx, text=text, ref=idx))
            idx += 1

        return ExtractionResult(
            segments=segments,
            structure={"segments": segments},
            warnings=list(dict.fromkeys(warnings)),
        )

    def reconstruct(self, original_path, translations, structure, output_path):
        document = docx.Document(str(original_path))
        segments = structure.get("segments", [])
        pointer = 0
        for para in iter_all_paragraphs(document):
            if not para.text.strip():
                continue
            if pointer >= len(segments):
                break
            seg = segments[pointer]
            pointer += 1
            tr = translations.get(seg.index)
            if tr is None:
                continue
            self._set_paragraph_text(para, tr)

        document.save(str(output_path))
        return output_path

    @staticmethod
    def _set_paragraph_text(para: Paragraph, text: str) -> None:
        runs = para.runs
        if not runs:
            para.add_run(text)
            return
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
