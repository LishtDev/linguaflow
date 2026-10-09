import openpyxl

from app.services.formatters.xlsx_formatter import XlsxFormatter


def _make_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feuil1"
    ws["A1"] = "Hello world"
    ws["B1"] = 42
    ws["C1"] = "=SUM(B1:B2)"
    ws["A2"] = "Note"
    wb.save(str(path))


def test_xlsx_translates_text_only(tmp_path):
    src = tmp_path / "book.xlsx"
    _make_xlsx(src)
    fmt = XlsxFormatter()
    result = fmt.extract(src)
    refs = [s.ref for s in result.segments]
    assert ("Feuil1", "A1") in refs
    assert ("Feuil1", "A2") in refs
    assert ("Feuil1", "B1") not in refs  # numeric
    assert ("Feuil1", "C1") not in refs  # formula

    translations = {s.index: "[FR]" + s.text for s in result.segments}
    out = tmp_path / "out.xlsx"
    fmt.reconstruct(src, translations, result.structure, out)

    wb = openpyxl.load_workbook(str(out))
    ws = wb["Feuil1"]
    assert ws["A1"].value == "[FR]Hello world"
    assert ws["B1"].value == 42
    assert ws["C1"].value == "=SUM(B1:B2)"  # formula preserved exactly
    assert ws.title == "Feuil1"
