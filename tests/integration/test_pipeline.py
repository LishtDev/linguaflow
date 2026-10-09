import hashlib
from pathlib import Path

import docx

from app.services.translation_service import TranslationPipeline


def _make_task(db, tmp_path, filename, content_bytes, ext, source_lang="en"):
    from app.services.task_service import TaskService
    from app.core.security import new_task_id
    src = tmp_path / filename
    src.write_bytes(content_bytes)
    task_id = new_task_id()
    ts = TaskService(db)
    ts.create(task_id=task_id, filename=filename, file_type=ext,
              source_path=str(src), target_lang="fr", source_lang=source_lang,
              checksum=hashlib.sha256(content_bytes).hexdigest())
    return task_id, src


def test_pipeline_txt_end_to_end(db, tmp_path, mock_engine):
    task_id, _ = _make_task(db, tmp_path, "a.txt", b"Hello world\nSecond line\n", "txt")
    pipeline = TranslationPipeline(db)
    result = pipeline.run(task_id)
    assert result["status"] == "completed"
    out = Path(result["result_path"])
    assert out.exists()
    text = out.read_text(encoding="utf-8")
    assert "bonjour le monde" in text.lower()
    assert "[FR]Second line" in text
    assert result["engine"] == "mock"


def test_pipeline_preserves_protected_elements(db, tmp_path, mock_engine):
    payload = "Contact admin@site.fr via https://example.com and use `run()`.\n"
    task_id, _ = _make_task(db, tmp_path, "b.txt", payload.encode(), "txt")
    result = TranslationPipeline(db).run(task_id)
    text = Path(result["result_path"]).read_text(encoding="utf-8")
    assert "admin@site.fr" in text
    assert "https://example.com" in text
    assert "run()" in text


def test_pipeline_docx_structure(db, tmp_path, mock_engine):
    src = tmp_path / "doc.docx"
    d = docx.Document()
    d.add_paragraph("Hello world")
    d.save(str(src))
    from app.services.task_service import TaskService
    from app.core.security import new_task_id
    task_id = new_task_id()
    TaskService(db).create(task_id=task_id, filename="doc.docx", file_type="docx",
                           source_path=str(src), target_lang="fr", source_lang="en")
    result = TranslationPipeline(db).run(task_id)
    rebuilt = docx.Document(result["result_path"])
    assert "bonjour le monde" in rebuilt.paragraphs[0].text.lower()


def test_pipeline_language_detection_recorded(db, tmp_path, mock_engine):
    from app.services.task_service import TaskService
    task_id, _ = _make_task(db, tmp_path, "c.txt",
                            b"This is an English sentence about testing.", "txt",
                            source_lang=None)
    TranslationPipeline(db).run(task_id)
    task = TaskService(db).get(task_id)
    assert task.detected_lang == "en" or task.source_lang == "en"
