import hashlib
from pathlib import Path

from app.core.security import new_task_id
from app.services.task_service import TaskService
from app.services.translation_service import TranslationPipeline


def _run(db, path, task_id=None):
    task_id = task_id or new_task_id()
    TaskService(db).create(task_id=task_id, filename=path.name,
                           file_type=path.suffix.lstrip("."),
                           source_path=str(path), target_lang="fr", source_lang="en",
                           checksum=hashlib.sha256(path.read_bytes()).hexdigest())
    return TranslationPipeline(db).run(task_id)


def test_pipeline_csv(db, tmp_path, mock_engine):
    f = tmp_path / "data.csv"
    f.write_text("name,price\nHello,1.5\nWorld,2\n", encoding="utf-8")
    result = _run(db, f)
    text = Path(result["result_path"]).read_text(encoding="utf-8")
    assert "bonjour" in text.lower()
    assert "1.5" in text


def test_pipeline_srt(db, tmp_path, mock_engine):
    f = tmp_path / "s.srt"
    f.write_text("1\n00:00:01,000 --> 00:00:04,000\nHello world\n", encoding="utf-8")
    result = _run(db, f)
    text = Path(result["result_path"]).read_text(encoding="utf-8")
    assert "00:00:01,000 --> 00:00:04,000" in text
    assert "bonjour le monde" in text.lower()


def test_pipeline_markdown(db, tmp_path, mock_engine):
    f = tmp_path / "d.md"
    f.write_text("# Hello world\n\n```\ncode here\n```\n", encoding="utf-8")
    result = _run(db, f)
    text = Path(result["result_path"]).read_text(encoding="utf-8")
    assert "# bonjour le monde" in text.lower()
    assert "code here" in text  # code block untouched


def test_pipeline_html(db, tmp_path, mock_engine):
    f = tmp_path / "p.html"
    f.write_text("<html><body><h1>Hello world</h1>"
                 "<script>var a=1;</script></body></html>", encoding="utf-8")
    result = _run(db, f)
    text = Path(result["result_path"]).read_text(encoding="utf-8")
    assert "bonjour le monde" in text.lower()
    assert "var a=1;" in text


def test_pipeline_html_keeps_doctype(db, tmp_path, mock_engine):
    f = tmp_path / "d.html"
    f.write_text("<!DOCTYPE html><html><body><p>Hello world</p></body></html>",
                 encoding="utf-8")
    result = _run(db, f)
    text = Path(result["result_path"]).read_text(encoding="utf-8")
    assert text.startswith("<!DOCTYPE html>")
    assert "bonjour le monde" in text.lower()


def test_pipeline_xlsx_formulas(db, tmp_path, mock_engine):
    import openpyxl
    f = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Hello world"
    ws["B1"] = '=A1&" x"'
    ws["A2"] = 42
    wb.save(str(f))
    result = _run(db, f)
    out = openpyxl.load_workbook(result["result_path"])
    ws = out.active
    assert ws["A1"].value == "bonjour le monde"
    assert ws["B1"].value == '=A1&" x"'
    assert ws["A2"].value == 42


def test_pipeline_pptx(db, tmp_path, mock_engine):
    from pptx import Presentation

    f = tmp_path / "p.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 0, 0)
    box.left, box.top, box.width, box.height = (
        914400, 914400, 5486400, 1828800)
    box.text_frame.text = "Hello world"
    prs.save(str(f))
    result = _run(db, f)

    out = Presentation(result["result_path"])
    slide = out.slides[0]
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert any("bonjour le monde" in t.lower() for t in texts)


def test_pipeline_pdf_textual(db, tmp_path, mock_engine):
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    f = tmp_path / "doc.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello world", fontsize=14)
    doc.save(str(f))
    doc.close()
    result = _run(db, f)
    out = fitz.open(result["result_path"])
    text = "".join(p.get_text() for p in out)
    out.close()
    assert "bonjour le monde" in text.lower()


def test_pipeline_epub(db, tmp_path, mock_engine):
    from ebooklib import epub

    f = tmp_path / "book.epub"
    book = epub.EpubBook()
    book.set_identifier("t1")
    book.set_title("T")
    book.set_language("en")
    ch = epub.EpubHtml(title="C", file_name="chap.xhtml", lang="en")
    ch.content = "<p>Hello world</p>"
    book.add_item(ch)
    book.toc = (ch,)
    book.spine = ["nav", ch]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    epub.write_epub(str(f), book)

    result = _run(db, f)
    out = epub.read_epub(result["result_path"])
    docs = [it.get_content().decode() for it in out.get_items()
            if it.get_name().endswith((".xhtml", ".html"))]
    assert "bonjour le monde" in "\n".join(docs).lower()
